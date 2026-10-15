from adm.services.permission_services import authorize_request
from rest_framework.views import APIView
from django.http import HttpResponse
from django.apps import apps
from openpyxl.styles import Font, PatternFill, Alignment
import openpyxl
from adm.services.query_services import exec_raw_sql, delete_exec_raw_sql
from rest_framework.views import APIView
from rest_framework import serializers
from rest_framework.response import Response
from rest_framework import status
from ..services.lead_services import *

from ..services.payment_services import *

# class DynamicExportExcel(APIView):

#     def get(self, request, app_label, model_name):

#         try:
#             # print(app_label)
#             # print(model_name)
#             # Model get pannrom
#             model = apps.get_model(app_label, model_name)
#         except LookupError:
#             return HttpResponse("Model not found", status=404)

#         # Workbook create
#         wb = openpyxl.Workbook()
#         ws = wb.active
#         ws.title = model_name

#         # Fields (only DB fields)
#         fields = [field.name for field in model._meta.fields]

#         # Header
#         ws.append(fields)

#         # Data fetch
#         queryset = model.objects.all()

#         for obj in queryset:
#             row = []
#             for field in fields:
#                 value = getattr(obj, field)

#                 # datetime format
#                 if hasattr(value, 'strftime'):
#                     value = value.strftime('%Y-%m-%d %H:%M:%S')

#                 row.append(str(value) if value is not None else '')

#             ws.append(row)

#         # Response
#         response = HttpResponse(
#             content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
#         )
#         filename = f"{model_name}.xlsx"
#         response['Content-Disposition'] = f'attachment; filename={filename}'

#         wb.save(response)
#         return response
    
    
    




from utils.excel_export import apply_excel_report_header, auto_fit_columns


class DynamicExportExcel(APIView):

    def post(self, request, *args, **kwargs):
        authorize_request('api_dynamic_export_excel', request.user)
        json_data = request.data.get('data')
        file_name = request.data.get('file_name', 'export')
        report_name = request.data.get('report_name') or str(file_name).replace('_', ' ').title()
        description = request.data.get('description') or f"All in one {report_name.lower()}"

        if not json_data:
            return HttpResponse("No data provided", status=400)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # ✅ Data extract and sheets create
        self.process_data(wb, json_data, report_name=report_name, description=description, user=request.user)

        # ✅ Sheet எதுவும் create ஆகலன்னா
        if not wb.sheetnames:
            ws = wb.create_sheet("Sheet1")
            apply_excel_report_header(ws, report_name=report_name, description=description, user=request.user)
            ws.cell(row=7, column=1, value="No Data")

        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename={file_name}.xlsx'
        wb.save(response)
        return response

    def process_data(self, wb, data, sheet_name="Sheet", report_name="Report", description="All in one report", user=None):
        """
        எந்த structure வந்தாலும் handle பண்ணும்:
        - list of dicts → direct sheet
        - dict with 'data' key → sheet
        - nested dict → recurse
        """

        # Case 1: List of flat dicts → direct sheet create
        if isinstance(data, list) and data and isinstance(data[0], dict):
            flat_rows = [row for row in data if not any(isinstance(v, (dict, list)) for v in row.values())]

            if flat_rows:
                self.create_sheet(wb, sheet_name, flat_rows, report_name=report_name, description=description, user=user)

            # Nested objects இருந்தா recurse
            for row in data:
                for key, value in row.items():
                    if isinstance(value, (dict, list)):
                        self.process_data(wb, value, key, report_name=f"{report_name} - {key.title()}", description=description, user=user)

        # Case 2: Dict with 'data' key → sheet create
        elif isinstance(data, dict) and 'data' in data:
            rows = data.get('data', [])
            count = data.get('count', None)
            if isinstance(rows, list) and rows:
                self.create_sheet(wb, sheet_name, rows, count, report_name=report_name, description=description, user=user)

        # Case 3: Nested dict → each key recurse
        elif isinstance(data, dict):
            for key, value in data.items():
                self.process_data(wb, value, key, report_name=f"{report_name} - {key.title()}", description=description, user=user)

    def create_sheet(self, wb, sheet_name, rows, count=None, report_name="Report", description="All in one report", user=None):
        if not rows:
            return

        clean_name = str(sheet_name).replace('_', ' ').title()[:31]
        base_name = clean_name
        counter = 1
        while clean_name in wb.sheetnames:
            clean_name = f"{base_name[:28]}_{counter}"
            counter += 1

        ws = wb.create_sheet(title=clean_name)

        # 1. Apply Top Metadata Header (Rows 1-4, spacers 5-6, returns row 7 for table headers)
        start_row = apply_excel_report_header(
            ws,
            report_name=report_name,
            description=description,
            user=user
        )

        header_font  = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill  = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")
        data_font    = Font(name="Calibri", size=10)
        border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # Optional count badge on Row 5 if provided
        if count is not None:
            count_cell = ws.cell(row=5, column=1, value=f"Total Records : {count}")
            count_cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
            count_cell.fill = PatternFill(start_color="E74C3C", end_color="E74C3C", fill_type="solid")
            count_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[5].height = 20

        # 2. Table Column Headers (Row 7)
        headers = list(rows[0].keys())
        ws.row_dimensions[start_row].height = 28
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=start_row, column=col_num, value=str(header).replace('_', ' ').upper())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border

        # 3. Data Rows (Row 8 onwards)
        current_row = start_row + 1
        for row_data in rows:
            ws.row_dimensions[current_row].height = 20
            for col_num, col in enumerate(headers, 1):
                val = row_data.get(col, "")
                cell = ws.cell(row=current_row, column=col_num, value=str(val) if val is not None else "")
                cell.font = data_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border = border
            current_row += 1

        # 4. Auto Column Widths
        auto_fit_columns(ws, start_row=start_row, min_width=14, col_a_min=18, col_b_min=28)

            
            







