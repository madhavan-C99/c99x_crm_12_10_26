import asyncio
import zoneinfo
from datetime import datetime, timedelta, date
from django.db.models import Count, Sum, Q, FloatField
from django.db.models.functions import TruncMonth, Coalesce
from django.utils import timezone
from channels.layers import get_channel_layer
from asgiref.sync import async_to_sync
from .query_services import exec_raw_sql
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import APIException

from telecalling.models import (
    Lead, CallDetails, FollowUp, PaymentInfo, 
    LeadSource, LossReason, PipelineStage, CampaignName
)
from adm.models import User, PipelineCategory


# ----------------- Date Filter Helper -----------------
def parse_date_filter(date_filter_str, custom_from=None, custom_to=None):
    try:
        ist = zoneinfo.ZoneInfo('Asia/Kolkata')
        today = datetime.now(ist).date()
    except Exception:
        today = datetime.now().date()

    raw_str = str(date_filter_str or "All").strip()

    # Priority 1: Direct custom_from / custom_to parameters
    if custom_from and custom_to:
        try:
            f_date = datetime.strptime(str(custom_from)[:10], '%Y-%m-%d').date() if isinstance(custom_from, str) else custom_from
            t_date = datetime.strptime(str(custom_to)[:10], '%Y-%m-%d').date() if isinstance(custom_to, str) else custom_to
            return f_date, t_date, f"{f_date.strftime('%b %d')} - {t_date.strftime('%b %d')}"
        except Exception:
            pass

    # Priority 2: String range format "YYYY-MM-DD to YYYY-MM-DD" or "YYYY-MM-DD - YYYY-MM-DD"
    if " to " in raw_str.lower() or " - " in raw_str:
        delimiter = " to " if " to " in raw_str.lower() else " - "
        parts = [p.strip() for p in raw_str.split(delimiter, 1)]
        if len(parts) == 2 and parts[0] and parts[1]:
            try:
                f_date = datetime.strptime(parts[0][:10], '%Y-%m-%d').date()
                t_date = datetime.strptime(parts[1][:10], '%Y-%m-%d').date()
                return f_date, t_date, f"{f_date.strftime('%b %d')} - {t_date.strftime('%b %d')}"
            except Exception:
                pass

    df = raw_str.lower().replace("-", "_")

    if df in ["all", "all time", "overall", "all_time", "none", ""]:
        return None, None, "overall"
    elif df in ["today", "daily"]:
        return today, today, "vs yesterday"
    elif df == "yesterday":
        y = today - timedelta(days=1)
        return y, y, "vs previous day"
    elif df in ["last_7_days", "last 7 days", "this week", "weekly", "week"]:
        start = today - timedelta(days=7)
        return start, today, "vs last 7 days"
    elif df in ["this_month", "this month", "monthly", "month"]:
        start = today.replace(day=1)
        return start, today, "vs last month"
    elif df in ["last_30_days", "last 30 days", "30 days"]:
        start = today - timedelta(days=30)
        return start, today, "vs previous 30 days"
    elif df in ["this_year", "this year", "yearly", "year"]:
        start = today.replace(month=1, day=1)
        return start, today, "vs last year"
    elif df in ["last_year", "last year", "previous_year", "previous year"]:
        start_ly = date(today.year - 1, 1, 1)
        end_ly = date(today.year - 1, 12, 31)
        return start_ly, end_ly, f"vs {today.year - 2}"
    else:
        return None, None, "overall"


def calc_metric_trend(current_val, prev_val, trend_label, is_percentage=False):
    if not trend_label or trend_label == "overall":
        return f"+0% {trend_label}" if trend_label else "+0% overall"
    if is_percentage:
        diff = round(current_val - prev_val)
        sign = "+" if diff > 0 else ""
        return f"{sign}{diff}% {trend_label}"
    if prev_val == 0:
        pct = 100 if current_val > 0 else 0
    else:
        pct = round(((current_val - prev_val) / prev_val) * 100)
    sign = "+" if pct > 0 else ""
    return f"{sign}{pct}% {trend_label}"


# ----------------- Master Data Payload Engine -----------------

def get_admin_dashboard_payload(user=None, pipeline_id=None, campaign_id=None, date_filter="All", from_date=None, to_date=None):
    from_date, to_date, trend_label = parse_date_filter(date_filter, from_date, to_date)

    # Resolve pipeline_id if string name (e.g. "Education") was passed
    if pipeline_id and not str(pipeline_id).isdigit():
        p_obj = PipelineCategory.objects.filter(
            Q(category_name__iexact=str(pipeline_id)) | Q(display_name__iexact=str(pipeline_id))
        ).first()
        if p_obj:
            pipeline_id = p_obj.id

    # 1. Base Queryset (Scoped to Pipeline, Campaign & Organization)
    base_leads_qs = Lead.objects.all()
    if user and getattr(user, 'is_authenticated', False) and getattr(user, 'organization', None):
        base_leads_qs = base_leads_qs.filter(Q(organization=user.organization) | Q(organization__isnull=True))
    if pipeline_id and str(pipeline_id).isdigit() and int(pipeline_id) > 0:
        base_leads_qs = base_leads_qs.filter(
            Q(pipeline_id=int(pipeline_id)) |
            Q(pipeline_stage__pipeline_category_id=int(pipeline_id)) |
            Q(campaign__pipeline_category_id=int(pipeline_id))
        )
    if campaign_id and str(campaign_id).isdigit() and int(campaign_id) > 0:
        base_leads_qs = base_leads_qs.filter(campaign_id=int(campaign_id))

    # 2. Date-Filtered Leads Queryset (For cards, funnel, and lead sources)
    if from_date and to_date:
        leads_qs = base_leads_qs.filter(created_at__date__range=[from_date, to_date])
    else:
        leads_qs = base_leads_qs  # 'All' selected -> overall pipeline leads

    # 1. METRICS
    total_leads_count = leads_qs.count()

    # Calls Made (Scoped to pipeline, campaign and date)
    calls_qs = CallDetails.objects.all()
    if user and getattr(user, 'is_authenticated', False) and getattr(user, 'organization', None):
        calls_qs = calls_qs.filter(Q(organization=user.organization) | Q(organization__isnull=True))
    if pipeline_id and str(pipeline_id).isdigit() and int(pipeline_id) > 0:
        calls_qs = calls_qs.filter(
            Q(lead__pipeline_id=int(pipeline_id)) |
            Q(lead__pipeline_stage__pipeline_category_id=int(pipeline_id)) |
            Q(lead__campaign__pipeline_category_id=int(pipeline_id))
        )
    if campaign_id and str(campaign_id).isdigit() and int(campaign_id) > 0:
        calls_qs = calls_qs.filter(lead__campaign_id=int(campaign_id))

    if from_date and to_date:
        calls_made = calls_qs.filter(created_at__date__range=[from_date, to_date]).count()
    else:
        calls_made = calls_qs.count()

    # Follow Up / Working Leads (Within selected period)
    follow_up_count = leads_qs.filter(
        Q(pipeline_stage__name__icontains="follow") |
        Q(pipeline_stage__name__icontains="todo") |
        Q(pipeline_stage__name__icontains="track")
    ).count()

    try:
        now_ist = datetime.now(zoneinfo.ZoneInfo('Asia/Kolkata'))
    except Exception:
        now_ist = timezone.now()

    # Missed Follow-ups count (Exclude leads already Won or Loss)
    missed_base_qs = FollowUp.objects.filter(is_attended=False).exclude(
        Q(lead__pipeline_stage__stage_type__in=['closed_won', 'closed_lost', 'terminal_joined', 'terminal_closed']) |
        (Q(lead__pipeline_stage__name__icontains="won") & ~Q(lead__pipeline_stage__name__icontains="loss") & ~Q(lead__pipeline_stage__name__icontains="lost")) |
        Q(lead__pipeline_stage__name__icontains="loss") |
        Q(lead__pipeline_stage__name__icontains="lost") |
        (Q(lead__pipeline_stage__name__icontains="convert") & ~Q(lead__pipeline_stage__name__icontains="loss") & ~Q(lead__pipeline_stage__name__icontains="lost")) |
        Q(lead__pipeline_stage__name__icontains="join") |
        Q(lead__pipeline_stage__name__icontains="admit") |
        Q(lead__pipeline_stage__name__icontains="closed")
    )
    if user and getattr(user, 'is_authenticated', False) and getattr(user, 'organization', None):
        missed_base_qs = missed_base_qs.filter(
            Q(lead__organization=user.organization) | Q(lead__organization__isnull=True)
        )
    if pipeline_id and str(pipeline_id).isdigit() and int(pipeline_id) > 0:
        missed_base_qs = missed_base_qs.filter(
            Q(lead__pipeline_id=int(pipeline_id)) |
            Q(lead__pipeline_stage__pipeline_category_id=int(pipeline_id)) |
            Q(lead__campaign__pipeline_category_id=int(pipeline_id))
        )
    if campaign_id and str(campaign_id).isdigit() and int(campaign_id) > 0:
        missed_base_qs = missed_base_qs.filter(lead__campaign_id=int(campaign_id))

    if from_date and to_date:
        missed_count = missed_base_qs.filter(
            scheduled_at__date__range=[from_date, to_date],
            scheduled_at__lte=now_ist
        ).count()
    else:
        missed_count = missed_base_qs.filter(scheduled_at__lte=now_ist).count()

    won_stage_filter = (
        Q(pipeline_stage__stage_type='closed_won') |
        (Q(pipeline_stage__name__icontains="won") & ~Q(pipeline_stage__name__icontains="loss") & ~Q(pipeline_stage__name__icontains="lost")) |
        (Q(pipeline_stage__name__icontains="convert") & ~Q(pipeline_stage__name__icontains="loss") & ~Q(pipeline_stage__name__icontains="lost"))
    )
    loss_stage_filter = (
        Q(pipeline_stage__stage_type='closed_lost') |
        Q(pipeline_stage__name__icontains="loss") |
        Q(pipeline_stage__name__icontains="lost")
    )
    converted_count = leads_qs.filter(won_stage_filter).count()
    loss_count = leads_qs.filter(loss_stage_filter).count()

    conv_rate = f"{round((converted_count / total_leads_count) * 100)}%" if total_leads_count > 0 else "0%"

    # Payments (STRICT - NO FALLBACK TO ALL-TIME! Multi-tenant Organization Scoped)
    payments_qs = PaymentInfo.objects.all()
    if user and getattr(user, 'is_authenticated', False) and getattr(user, 'organization', None):
        payments_qs = payments_qs.filter(
            Q(lead__organization=user.organization) | Q(lead__organization__isnull=True)
        )
    if pipeline_id and str(pipeline_id).isdigit() and int(pipeline_id) > 0:
        payments_qs = payments_qs.filter(
            Q(lead__pipeline_id=int(pipeline_id)) |
            Q(lead__pipeline_stage__pipeline_category_id=int(pipeline_id)) |
            Q(lead__campaign__pipeline_category_id=int(pipeline_id))
        )
    if campaign_id and str(campaign_id).isdigit() and int(campaign_id) > 0:
        payments_qs = payments_qs.filter(lead__campaign_id=int(campaign_id))

    if from_date and to_date:
        period_payments = payments_qs.filter(created_at__date__range=[from_date, to_date])
        rev_sum = period_payments.aggregate(total=Sum('amount_paid'))['total'] or 0
        pending_sum = period_payments.aggregate(total=Sum('pending_amount'))['total'] or 0
    else:
        rev_sum = payments_qs.aggregate(total=Sum('amount_paid'))['total'] or 0
        pending_sum = payments_qs.aggregate(total=Sum('pending_amount'))['total'] or 0

    total_rev_str = f"{int(rev_sum):,}"
    pending_str = f"{int(pending_sum):,}"

    # Dynamic Trends Calculation (Compare with preceding period of identical duration)
    prev_from_date, prev_to_date = None, None
    if from_date and to_date:
        delta_days = (to_date - from_date).days + 1
        prev_to_date = from_date - timedelta(days=1)
        prev_from_date = prev_to_date - timedelta(days=delta_days - 1)

    if prev_from_date and prev_to_date:
        prev_leads_qs = base_leads_qs.filter(created_at__date__range=[prev_from_date, prev_to_date])
        prev_total_leads = prev_leads_qs.count()
        prev_calls = calls_qs.filter(created_at__date__range=[prev_from_date, prev_to_date]).count()
        prev_follow_up = prev_leads_qs.filter(
            Q(pipeline_stage__name__icontains="follow") |
            Q(pipeline_stage__name__icontains="todo") |
            Q(pipeline_stage__name__icontains="track")
        ).count()
        prev_missed = missed_base_qs.filter(scheduled_at__date__range=[prev_from_date, prev_to_date]).count()
        prev_converted = prev_leads_qs.filter(won_stage_filter).count()
        prev_loss = prev_leads_qs.filter(loss_stage_filter).count()
        prev_conv_rate = (prev_converted / prev_total_leads) * 100 if prev_total_leads > 0 else 0

        prev_period_payments = payments_qs.filter(created_at__date__range=[prev_from_date, prev_to_date])
        prev_pending = prev_period_payments.aggregate(total=Sum('pending_amount'))['total'] or 0
    else:
        prev_total_leads = prev_calls = prev_follow_up = prev_missed = 0
        prev_converted = prev_loss = prev_conv_rate = prev_pending = 0

    num_conv_rate = (converted_count / total_leads_count) * 100 if total_leads_count > 0 else 0

    metrics = {
        "total_leads":      {"value": str(total_leads_count), "trend": calc_metric_trend(total_leads_count, prev_total_leads, trend_label)},
        "calls_made":       {"value": str(calls_made),        "trend": calc_metric_trend(calls_made, prev_calls, trend_label)},
        "follow_up":        {"value": str(follow_up_count),   "trend": calc_metric_trend(follow_up_count, prev_follow_up, trend_label)},
        "missed_follow_up": {"value": str(missed_count),      "trend": calc_metric_trend(missed_count, prev_missed, trend_label)},
        "converted_leads":  {"value": str(converted_count),   "trend": calc_metric_trend(converted_count, prev_converted, trend_label)},
        "loss_leads":       {"value": str(loss_count),        "trend": calc_metric_trend(loss_count, prev_loss, trend_label)},
        "conversion_rate":  {"value": conv_rate.replace("%", ""), "trend": calc_metric_trend(num_conv_rate, prev_conv_rate, trend_label, is_percentage=True)},
        "pending_payment":  {"value": pending_str,            "trend": calc_metric_trend(pending_sum, prev_pending, trend_label)}
    }

    # 2. MISSED FOLLOW-UPS (Table rows)
    missed_table_qs = missed_base_qs.filter(scheduled_at__lte=now_ist)
    if from_date and to_date:
        missed_table_qs = missed_table_qs.filter(scheduled_at__date__range=[from_date, to_date])

    missed_records = missed_table_qs.select_related('lead', 'telecaller', 'lead__pipeline_stage').order_by('-scheduled_at')[:20]

    missed_followups = []
    for mf in missed_records:
        sched_date = mf.scheduled_at.strftime('%b %d') if mf.scheduled_at else "Today"
        is_urgent = False
        if mf.scheduled_at:
            try:
                is_urgent = (now_ist.date() - mf.scheduled_at.date()).days >= 1
            except Exception:
                is_urgent = False

        telecaller_name = "Unassigned"
        if mf.telecaller:
            telecaller_name = mf.telecaller.get_full_name() or mf.telecaller.username

        mob = (mf.lead.mobile_no or mf.lead.alternative_mobile or "-") if mf.lead else "-"

        missed_followups.append({
            "id": mf.id,
            "followup_id": mf.id,
            "lead_id": mf.lead_id,
            "student_name": mf.lead.full_name if mf.lead else "Unknown",
            "mobile_no": mob,
            "mobile": mob,
            "mobile_number": mob,
            "phone": mob,
            "contact_no": mob,
            "lead_stage": mf.lead.pipeline_stage.name if (mf.lead and mf.lead.pipeline_stage) else "Hot",
            "follow_up_date": sched_date,
            "telecaller": telecaller_name,
            "flag": "red" if is_urgent else "orange"
        })

    # 3. PIPELINE STAGES (Funnel - from filtered leads_qs)
    stages = PipelineStage.objects.filter(is_active=True).order_by('order_no')
    if pipeline_id and str(pipeline_id).isdigit() and int(pipeline_id) > 0:
        stages = stages.filter(pipeline_category_id=int(pipeline_id))

    stage_counts = leads_qs.values('pipeline_stage_id').annotate(count=Count('id'))
    counts_map = {item['pipeline_stage_id']: item['count'] for item in stage_counts}

    pipeline_stages = []
    for stage in stages:
        pipeline_stages.append({
            "stage_id": stage.id,
            "stage_name": stage.name,
            "stage_slug": stage.name.lower().replace(" ", "_"),
            "count": counts_map.get(stage.id, 0),
            "stage_type": stage.stage_type
        })

    # 4. LEAD SOURCES (Distribution - from filtered leads_qs)
    sources = LeadSource.objects.filter(is_active=True)
    source_counts = leads_qs.values('lead_source_id').annotate(count=Count('id'))
    src_counts_map = {item['lead_source_id']: item['count'] for item in source_counts}

    lead_sources = []
    for src in sources:
        cnt = src_counts_map.get(src.id, 0)
        pct = f"{round((cnt / total_leads_count) * 100)}%" if total_leads_count > 0 else "0%"
        lead_sources.append({
            "source_id": src.id,
            "source_name": src.name,
            "count": cnt,
            "percentage": pct
        })

    # 5. CHARTS (Lead Status By Month & Revenue By Month)
    twelve_months_ago = now_ist.date().replace(day=1) - timedelta(days=365)
    
    monthly_leads_qs = base_leads_qs.filter(created_at__date__gte=twelve_months_ago)
    if from_date and to_date:
        monthly_leads_qs = monthly_leads_qs.filter(created_at__date__range=[from_date, to_date])

    monthly_stats = monthly_leads_qs.annotate(
        month=TruncMonth('created_at')
    ).values('month').annotate(
        total=Count('id'),
        won=Count('id', filter=won_stage_filter),
        loss=Count('id', filter=loss_stage_filter)
    ).order_by('month')

    months_list = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    leads_status_chart = {
        "labels": [],
        "datasets": {
            "total_leads": [],
            "won_leads": [],
            "loss_leads": []
        }
    }
    for stat in monthly_stats:
        if stat.get('month'):
            m_label = months_list[stat['month'].month - 1]
            leads_status_chart["labels"].append(m_label)
            leads_status_chart["datasets"]["total_leads"].append(stat['total'])
            leads_status_chart["datasets"]["won_leads"].append(stat['won'])
            leads_status_chart["datasets"]["loss_leads"].append(stat['loss'])

    # Revenue by Month
    monthly_pay_qs = payments_qs.filter(created_at__date__gte=twelve_months_ago)
    if from_date and to_date:
        monthly_pay_qs = monthly_pay_qs.filter(created_at__date__range=[from_date, to_date])

    rev_stats = monthly_pay_qs.annotate(
        month=TruncMonth('created_at')
    ).values('month').annotate(
        revenue=Coalesce(Sum('amount_paid'), 0.0, output_field=FloatField())
    ).order_by('month')

    revenue_chart = {
        "labels": [],
        "data": []
    }
    for rstat in rev_stats:
        if rstat.get('month'):
            m_label = months_list[rstat['month'].month - 1]
            revenue_chart["labels"].append(m_label)
            revenue_chart["data"].append(int(rstat['revenue']))

    # 6. LOSS REASON ANALYSIS
    loss_leads_qs = leads_qs.filter(loss_stage_filter)
    loss_counts = loss_leads_qs.values('loss_reason_id').annotate(count=Count('id')).order_by('-count')
    reason_map = {r.id: r.name for r in LossReason.objects.filter(is_active=True)}
    loss_reasons = []
    for lc in loss_counts:
        rid = lc.get('loss_reason_id')
        if rid and rid in reason_map:
            cnt = lc['count']
            pct = f"{round((cnt / loss_count) * 100)}%" if loss_count > 0 else "0%"
            loss_reasons.append({
                "reason": reason_map[rid],
                "count": cnt,
                "percentage": pct
            })

    return {
        "metrics": metrics,
        "missed_followups": missed_followups,
        "pipeline_stages": pipeline_stages,
        "lead_sources": lead_sources,
        "lead_status_chart": leads_status_chart,
        "revenue_chart": revenue_chart,
        "loss_reason_analysis": loss_reasons
    }


# ----------------- Decoupled Live Trigger (TL Required Layer) -----------------

def trigger_dashboard_update():
    try:
        channel_layer = get_channel_layer()
        if not channel_layer:
            return

        message = {
            "type": "broadcast_dashboard_update"
        }

        # Check if an event loop is already active in the current thread
        try:
            loop = asyncio.get_running_loop()
        except RuntimeError:
            loop = None

        if loop and loop.is_running():
            # In Async Context (ASGI / Async views / Async test scripts)
            asyncio.ensure_future(
                channel_layer.group_send("admin_dashboard", message),
                loop=loop
            )
        else:
            # In Synchronous Thread (Standard Django views / Huey workers)
            async_to_sync(channel_layer.group_send)("admin_dashboard", message)

    except Exception as e:
        print(f"[WS TRIGGER ERROR] {e}")


def normalize_dashboard_date_filter(date_filter="Today", from_date=None, to_date=None):
    today = timezone.localdate()
    df = str(date_filter or "Today").strip().lower()

    if df in ("today", "daily"):
        return today.isoformat(), today.isoformat()
    elif df == "yesterday":
        yest = today - timedelta(days=1)
        return yest.isoformat(), yest.isoformat()
    elif df in ("last_7_days", "last 7 days", "this week", "weekly", "week"):
        start_w = today - timedelta(days=7)
        return start_w.isoformat(), today.isoformat()
    elif df in ("this_month", "this month", "monthly", "month"):
        start_m = today.replace(day=1)
        return start_m.isoformat(), today.isoformat()
    elif df in ("last_30_days", "last 30 days", "30 days"):
        start_30 = today - timedelta(days=30)
        return start_30.isoformat(), today.isoformat()
    elif df in ("yearly", "year", "this_year"):
        start_y = date(today.year, 1, 1)
        return start_y.isoformat(), today.isoformat()
    elif df == "custom" and from_date and to_date:
        return str(from_date), str(to_date)
    elif df in ("all", "all_time", "overall"):
        return "", ""
    
    return today.isoformat(), today.isoformat()


def fetch_dashboard_pending_payments(user=None, pipeline_id=None, campaign_id=None, date_filter="Today", from_date=None, to_date=None, **kwargs):
    org_id = getattr(user, 'organization_id', 0) if user else 0
    from_d, to_d = normalize_dashboard_date_filter(date_filter, from_date, to_date)

    qry_vars = {
        "organization_id": org_id,
        "pipeline_id": pipeline_id if pipeline_id not in (None, "", "null") else 0,
        "campaign_id": campaign_id if campaign_id not in (None, "", "null") else 0,
        "from_date": from_d,
        "to_date": to_d,
    }

    records = exec_raw_sql('D_ADMIN_DASHBOARD_PENDING_PAYMENTS', qry_vars)
    return records or []


def fetch_dashboard_performance_overview(user=None, pipeline_id=None, campaign_id=None, date_filter="Today", from_date=None, to_date=None, **kwargs):
    org_id = getattr(user, 'organization_id', 0) if user else 0
    from_d, to_d = normalize_dashboard_date_filter(date_filter, from_date, to_date)

    qry_vars = {
        "organization_id": org_id,
        "pipeline_id": pipeline_id if pipeline_id not in (None, "", "null") else 0,
        "campaign_id": campaign_id if campaign_id not in (None, "", "null") else 0,
        "from_date": from_d,
        "to_date": to_d,
    }

    records = exec_raw_sql('D_ADMIN_DASHBOARD_PERFORMANCE_OVERVIEW', qry_vars)
    return records or []


def fetch_dashboard_best_performer(user=None, pipeline_id=None, campaign_id=None, date_filter="Today", from_date=None, to_date=None, **kwargs):
    org_id = getattr(user, 'organization_id', 0) if user else 0
    from_d, to_d = normalize_dashboard_date_filter(date_filter, from_date, to_date)

    qry_vars = {
        "organization_id": org_id,
        "pipeline_id": pipeline_id if pipeline_id not in (None, "", "null") else 0,
        "campaign_id": campaign_id if campaign_id not in (None, "", "null") else 0,
        "from_date": from_d,
        "to_date": to_d,
    }

    records = exec_raw_sql('D_ADMIN_DASHBOARD_BEST_PERFORMER', qry_vars)
    if records and len(records) > 0:
        best = records[0]
    else:
        best = {
            "name": "-",
            "avatar": "https://api.dicebear.com/7.x/avataaars/svg?seed=Admin",
            "deals_closed": 0,
            "calls_made": 0,
            "revenue": "₹0"
        }

    return best


def fetch_admin_dashboard_export_data(user=None, data: dict = None, **filters):
    raw = {}
    if data and isinstance(data, dict):
        raw.update(data)
    if filters:
        if "filters" in filters and isinstance(filters["filters"], dict):
            raw.update(filters.pop("filters"))
        raw.update(filters)

    pipeline_id = raw.get("pipeline_id")
    campaign_id = raw.get("campaign_id")
    date_filter = raw.get("date_filter", "This Year")
    from_date = raw.get("from_date")
    to_date = raw.get("to_date")

    # 1. Main Dashboard Payload (Tiles + Missed Follow Up Table)
    dashboard_main = get_admin_dashboard_payload(user, pipeline_id, campaign_id, date_filter, from_date, to_date)

    # 2. Payments & Performance (Tables 2 & 3)
    payments = fetch_dashboard_pending_payments(user, pipeline_id, campaign_id, date_filter, from_date, to_date)
    performance = fetch_dashboard_performance_overview(user, pipeline_id, campaign_id, date_filter, from_date, to_date)

    payments_list = payments if isinstance(payments, list) else (payments.get("data", []) if isinstance(payments, dict) else [])
    perf_list = performance if isinstance(performance, list) else (performance.get("data", []) if isinstance(performance, dict) else [])
    missed_list = dashboard_main.get("missed_followups", [])

    return {
        "generated_for": getattr(user, 'username', 'Admin'),
        "role": "Admin",
        "date_filter_label": date_filter,
        "metrics": dashboard_main.get("metrics", {}),       # 📌 8 TILES
        "missed_followups": missed_list,                    # 📌 TABLE 1: Missed Follow Up
        "pending_payments": payments_list,                  # 📌 TABLE 2: Pending Payment List
        "performance_overview": perf_list,                  # 📌 TABLE 3: Performance Overview
    }


def generate_admin_dashboard_excel(file_path, dashboard_data, user=None):
    wb = openpyxl.Workbook()

    header_fill = PatternFill(fill_type="solid", fgColor="16A34A")  # Brand Green
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    title_font = Font(name="Calibri", size=14, bold=True, color="1E293B")
    meta_font = Font(name="Calibri", size=9, italic=True, color="64748B")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    # ─────────────────────────────────────────────────────────────
    # 📑 SHEET 1: TILES (8 KPI Metrics Summary)
    # ─────────────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "KPI Summary"
    ws1.cell(row=1, column=1, value="Admin Dashboard - KPI Summary").font = title_font
    ws1.cell(row=2, column=1, value=f"Generated on {timezone.now().strftime('%d %b %Y, %I:%M %p')} | Filter: {dashboard_data.get('date_filter_label', 'This Year')}").font = meta_font

    kpi_headers = ["S.No", "Metric Name (Tile)", "Current Count / Value", "Trend vs Last Period"]
    for col_idx, h in enumerate(kpi_headers, start=1):
        c = ws1.cell(row=4, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align
        c.border = border

    metrics = dashboard_data.get("metrics", {})
    tiles_data = [
        ("Total Leads", metrics.get("total_leads", {}).get("value", 0), metrics.get("total_leads", {}).get("trend", "")),
        ("Calls Made", metrics.get("calls_made", {}).get("value", 0), metrics.get("calls_made", {}).get("trend", "")),
        ("Follow - Up", metrics.get("follow_up", {}).get("value", 0), metrics.get("follow_up", {}).get("trend", "")),
        ("Missed Follow Up", metrics.get("missed_follow_up", {}).get("value", 0), metrics.get("missed_follow_up", {}).get("trend", "")),
        ("Converted Leads", metrics.get("converted_leads", {}).get("value", 0), metrics.get("converted_leads", {}).get("trend", "")),
        ("Loss Leads", metrics.get("loss_leads", {}).get("value", 0), metrics.get("loss_leads", {}).get("trend", "")),
        ("Conversion Rate", f"{metrics.get('conversion_rate', {}).get('value', 0)}%", metrics.get("conversion_rate", {}).get("trend", "")),
        ("Pending Payment", f"Rs. {metrics.get('pending_payment', {}).get('value', 0)}", metrics.get("pending_payment", {}).get("trend", ""))
    ]

    for idx, (label, val, trend) in enumerate(tiles_data, start=1):
        r_i = idx + 4
        ws1.row_dimensions[r_i].height = 22
        vals = [idx, label, val, trend]
        for c_i, v in enumerate(vals, start=1):
            cell = ws1.cell(row=r_i, column=c_i, value=v)
            cell.font = data_font
            cell.alignment = center_align if c_i != 2 else left_align
            cell.border = border

    # ─────────────────────────────────────────────────────────────
    # 📑 SHEET 2: TABLE 1 - Missed Follow Up
    # ─────────────────────────────────────────────────────────────
    ws2 = wb.create_sheet(title="Missed Follow Up")
    ws2.cell(row=1, column=1, value="Admin Dashboard - Missed Follow Up List").font = title_font
    ws2.cell(row=2, column=1, value=f"Total Records: {len(dashboard_data.get('missed_followups', []))}").font = meta_font

    missed_headers = ["S.No", "Student Name", "Mobile Number", "Lead Stage", "Follow up date", "Telecaller", "Urgency Flag"]
    for col_idx, h in enumerate(missed_headers, start=1):
        c = ws2.cell(row=4, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align
        c.border = border

    for idx, row in enumerate(dashboard_data.get("missed_followups", []), start=1):
        r_i = idx + 4
        ws2.row_dimensions[r_i].height = 20
        vals = [
            idx,
            row.get("student_name") or "-",
            row.get("mobile_number") or row.get("mobile_no") or "-",
            row.get("lead_stage") or "Follow Up",
            row.get("follow_up_date") or "-",
            row.get("telecaller") or "Unassigned",
            "Urgent (Red Flag)" if row.get("flag") == "red" else "Normal"
        ]
        for c_i, v in enumerate(vals, start=1):
            cell = ws2.cell(row=r_i, column=c_i, value=v)
            cell.font = data_font
            cell.alignment = left_align if c_i in [2, 6] else center_align
            cell.border = border

    # ─────────────────────────────────────────────────────────────
    # 📑 SHEET 3: TABLE 2 - Pending Payment List
    # ─────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet(title="Pending Payment List")
    ws3.cell(row=1, column=1, value="Admin Dashboard - Pending Payment List").font = title_font
    ws3.cell(row=2, column=1, value=f"Total Records: {len(dashboard_data.get('pending_payments', []))}").font = meta_font

    pay_headers = ["S.No", "Student Name", "Course", "Amount (Rs.)", "Due Date", "Status", "Telecaller"]
    for col_idx, h in enumerate(pay_headers, start=1):
        c = ws3.cell(row=4, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align
        c.border = border

    for idx, row in enumerate(dashboard_data.get("pending_payments", []), start=1):
        r_i = idx + 4
        ws3.row_dimensions[r_i].height = 20
        vals = [
            idx,
            row.get("student_name") or "-",
            row.get("course") or "-",
            row.get("amount") or 0,
            row.get("due_date") or "-",
            row.get("status") or "Pending",
            row.get("telecaller") or "-"
        ]
        for c_i, v in enumerate(vals, start=1):
            cell = ws3.cell(row=r_i, column=c_i, value=v)
            cell.font = data_font
            cell.alignment = left_align if c_i in [2, 3, 7] else center_align
            cell.border = border

    # ─────────────────────────────────────────────────────────────
    # 📑 SHEET 4: TABLE 3 - Performance Overview
    # ─────────────────────────────────────────────────────────────
    ws4 = wb.create_sheet(title="Performance Overview")
    ws4.cell(row=1, column=1, value="Admin Dashboard - Telecaller Performance Overview").font = title_font
    ws4.cell(row=2, column=1, value=f"Total Telecallers: {len(dashboard_data.get('performance_overview', []))}").font = meta_font

    perf_headers = ["S.No", "Telecaller", "Total Leads", "Calls Made", "Connected", "Missed Followup", "Conversion", "Rating"]
    for col_idx, h in enumerate(perf_headers, start=1):
        c = ws4.cell(row=4, column=col_idx, value=h)
        c.fill = header_fill
        c.font = header_font
        c.alignment = center_align
        c.border = border

    for idx, row in enumerate(dashboard_data.get("performance_overview", []), start=1):
        r_i = idx + 4
        ws4.row_dimensions[r_i].height = 20
        vals = [
            idx,
            row.get("telecaller") or "-",
            row.get("total_leads") or 0,
            row.get("calls_made") or 0,
            row.get("connected") or 0,
            row.get("missed_followup") or 0,
            row.get("conversion") or 0,
            f"{row.get('rating')}%" if str(row.get('rating', '')).isdigit() else str(row.get('rating') or '0%')
        ]
        for c_i, v in enumerate(vals, start=1):
            cell = ws4.cell(row=r_i, column=c_i, value=v)
            cell.font = data_font
            cell.alignment = left_align if c_i == 2 else center_align
            cell.border = border

    for sheet in [ws1, ws2, ws3, ws4]:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            sheet.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 4, 15)

    wb.save(file_path)
