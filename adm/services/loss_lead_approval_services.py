import os
import openpyxl
from datetime import datetime, date, timedelta
from django.conf import settings
from django.db.models import Q, Count
from django.utils import timezone
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.db import transaction
from rest_framework.exceptions import APIException, ValidationError, PermissionDenied

from telecalling.models import (
    Lead, LossLeadDetail, CallDetails, FollowUp, PipelineStage,
    LeadSource, CampaignName, CoursePlan, CourseName, User, LossReason, Notification
)
from ..models import AdminLossActionLog, AdminApprovedLossLead, PipelineCategory
from .lead_services import get_user_display_name

from telecalling.tasks import send_loss_request_result_notification,send_lead_assigned_notification


def fetch_loss_lead_approval_requests_admin(user=None, **data):
    
    try:
        today = timezone.now().date()
        
        approved_lead_ids = AdminApprovedLossLead.objects.values_list("lead_id", flat=True)
        
        base_qs = Lead.objects.filter(
            Q(pipeline_stage__name__icontains="loss") |
            Q(pipeline_stage__name__icontains="lost") |
            Q(current_status__iexact="loss") |
            Q(current_status__iexact="lost")
        ).select_related(
            "assigned_to", "pipeline_stage", "campaign", "lead_source", "course_plan", "course_name"
        ).distinct().order_by("-updated_at", "-created_at")

        if user and getattr(user, 'is_authenticated', False) and getattr(user, 'organization', None):
            base_qs = base_qs.filter(organization=user.organization)
        elif user and getattr(user, 'is_authenticated', False):
            base_qs = base_qs.none()

        # 2. Search Filter (name, phone, email)
        search = data.get("search")
        if search and str(search).strip() != "":
            search_clean = str(search).strip()
            base_qs = base_qs.filter(
                Q(full_name__icontains=search_clean) |
                Q(mobile_no__icontains=search_clean) |
                Q(email__icontains=search_clean)
            )

        # 3. Date Filter (Today, Yesterday, Weekly, Monthly, Yearly, Custom)
        date_filter = str(data.get("date_filter") or data.get("date_filter_type") or "all").lower().strip()
        from_date = data.get("from_date")
        to_date = data.get("to_date")

        filter_from = None
        filter_to = None

        if date_filter == "today":
            filter_from = today
            filter_to = today
        elif date_filter == "yesterday":
            filter_from = today - timedelta(days=1)
            filter_to = today - timedelta(days=1)
        elif date_filter in ["last_7_days", "last 7 days", "this_week", "weekly", "week"]:
            filter_from = today - timedelta(days=7)
            filter_to = today
        elif date_filter in ["last_30_days", "last 30 days", "30 days"]:
            filter_from = today - timedelta(days=30)
            filter_to = today
        elif date_filter in ["this_month", "this month", "monthly", "month"]:
            filter_from = today.replace(day=1)
            filter_to = today
        elif date_filter in ["this_year", "yearly", "year"]:
            filter_from = date(today.year, 1, 1)
            filter_to = date(today.year, 12, 31)
        elif (date_filter == "custom" or from_date or to_date) and (from_date or to_date):
            if from_date:
                try:
                    filter_from = datetime.strptime(str(from_date)[:10], "%Y-%m-%d").date() if isinstance(from_date, str) else from_date
                except Exception:
                    filter_from = None
            if to_date:
                try:
                    filter_to = datetime.strptime(str(to_date)[:10], "%Y-%m-%d").date() if isinstance(to_date, str) else to_date
                except Exception:
                    filter_to = None

        if filter_from and filter_to:
            base_qs = base_qs.filter(
                Q(loss_detail__created_at__date__range=[filter_from, filter_to]) |
                Q(updated_at__date__range=[filter_from, filter_to]) |
                Q(enquiry_date__date__range=[filter_from, filter_to])
            ).distinct()
        elif filter_from:
            base_qs = base_qs.filter(
                Q(loss_detail__created_at__date__gte=filter_from) |
                Q(updated_at__date__gte=filter_from) |
                Q(enquiry_date__date__gte=filter_from)
            ).distinct()
        elif filter_to:
            base_qs = base_qs.filter(
                Q(loss_detail__created_at__date__lte=filter_to) |
                Q(updated_at__date__lte=filter_to) |
                Q(enquiry_date__date__lte=filter_to)
            ).distinct()

        # 4. Additional DB-Driven Dropdown & Popup Filters
        pipeline_id_val = data.get("pipeline_id")
        if pipeline_id_val and str(pipeline_id_val).isdigit() and int(pipeline_id_val) > 0:
            base_qs = base_qs.filter(
                Q(pipeline_stage__pipeline_category_id=int(pipeline_id_val)) |
                Q(campaign__pipeline_category_id=int(pipeline_id_val)) |
                Q(campaign_id=int(pipeline_id_val))
            )

        loss_reason_val = data.get("loss_reason") or data.get("loss_reason_id") or data.get("reason_id")
        if loss_reason_val and str(loss_reason_val).strip() and str(loss_reason_val).lower() not in ["all", "none"]:
            loss_reason_str = str(loss_reason_val).strip()
            if loss_reason_str.isdigit():
                r_id = int(loss_reason_str)
                base_qs = base_qs.filter(
                    Q(loss_reason_id=r_id) |
                    Q(loss_detail__loss_reason_id=r_id) |
                    Q(loss_detail__main_reason_id=r_id)
                )
            else:
                base_qs = base_qs.filter(
                    Q(loss_reason__name__icontains=loss_reason_str) |
                    Q(loss_detail__loss_reason__name__icontains=loss_reason_str) |
                    Q(loss_detail__detailed_reason__icontains=loss_reason_str)
                )

        telecaller_val = data.get("telecaller") or data.get("assigned_to_id") or data.get("telecaller_id") or data.get("tele_id")
        if telecaller_val and str(telecaller_val).strip() and str(telecaller_val).lower() not in ["all", "none"]:
            telecaller_str = str(telecaller_val).strip()
            if telecaller_str.isdigit():
                base_qs = base_qs.filter(assigned_to_id=int(telecaller_str))
            else:
                base_qs = base_qs.filter(
                    Q(assigned_to__username__iexact=telecaller_str) |
                    Q(assigned_to__first_name__icontains=telecaller_str) |
                    Q(assigned_to__last_name__icontains=telecaller_str)
                )

        course_val = data.get("course") or data.get("course_id") or data.get("course_name_id")
        if course_val and str(course_val).strip() and str(course_val).lower() not in ["all", "none"]:
            course_str = str(course_val).strip()
            if course_str.isdigit():
                base_qs = base_qs.filter(course_name_id=int(course_str))
            else:
                base_qs = base_qs.filter(
                    Q(course_name__coursename__icontains=course_str) |
                    Q(course_name__name__icontains=course_str)
                )

        lead_source_val = data.get("lead_source") or data.get("lead_source_id") or data.get("source_id")
        if lead_source_val and str(lead_source_val).strip() and str(lead_source_val).lower() not in ["all", "none"]:
            source_str = str(lead_source_val).strip()
            if source_str.isdigit():
                base_qs = base_qs.filter(lead_source_id=int(source_str))
            else:
                base_qs = base_qs.filter(lead_source__name__icontains=source_str)

        course_plan_val = data.get("course_plan_id")
        if course_plan_val and str(course_plan_val).isdigit():
            base_qs = base_qs.filter(course_plan_id=int(course_plan_val))

        campaign_val = data.get("campaign_id")
        if campaign_val and str(campaign_val).isdigit():
            base_qs = base_qs.filter(campaign_id=int(campaign_val))

        approval_status_val = str(data.get("approval_status") or data.get("status") or "all").lower().strip()
        if approval_status_val in ["approved", "approve"]:
            base_qs = base_qs.filter(approved_loss_record__isnull=False)
        elif approval_status_val in ["rejected", "reject"]:
            base_qs = base_qs.filter(admin_loss_action_logs__action_type='rejected').distinct()
        elif approval_status_val in ["reassigned", "reassign"]:
            base_qs = base_qs.filter(admin_loss_action_logs__action_type='reassigned').distinct()
        else:
            base_qs = base_qs.exclude(id__in=approved_lead_ids)

        # Sorting (sort_type: newest / oldest)
        sort_type = str(data.get("sort_type") or data.get("sort_by") or "newest").lower().strip()
        if sort_type in ["oldest", "asc", "created_at"]:
            base_qs = base_qs.order_by("updated_at", "created_at")
        else:
            base_qs = base_qs.order_by("-updated_at", "-created_at")

        total_count = base_qs.count()

        # 5. Pagination
        page_size_input = data.get("page_size")
        if str(page_size_input).lower() in ["0", "all", "none"] or page_size_input == 0:
            page = 1
            page_size = total_count if total_count > 0 else 1000
            start = 0
            rows = base_qs
        else:
            page = int(data.get("page") or 1)
            if page < 1:
                page = 1
            page_size = int(page_size_input or 50)
            start = (page - 1) * page_size
            end = start + page_size
            rows = base_qs[start:end]

        # 6. Populate Table Rows matching 11 Figma Columns
        leads_list = []
        for idx, lead in enumerate(rows, start=start + 1):
            calls = CallDetails.objects.filter(lead=lead).order_by("-created_at")
            total_calls = calls.count()
            latest_call = calls.first()
            loss_detail = LossLeadDetail.objects.filter(lead=lead).first()

            # Effort Summary (e.g. "12 Calls Done")
            effort_summary = f"{total_calls} Calls Done" if total_calls > 0 else "0 Calls Done"

            # Loss Reason
            loss_reason_str = "-"
            if loss_detail:
                reason_obj = loss_detail.loss_reason
                if reason_obj:
                    loss_reason_str = reason_obj.name
                elif loss_detail.detailed_reason:
                    loss_reason_str = loss_detail.detailed_reason

            # Last Conversation Outcome (actual call disposition / outcome / summary)
            last_conversation_outcome = "-"
            if latest_call:
                call_tag = None
                if getattr(latest_call, 'select_tag', None):
                    call_tag = getattr(latest_call.select_tag, 'display_value', None) or latest_call.select_tag.name
                elif getattr(latest_call, 'stage', None):
                    call_tag = getattr(latest_call.stage, 'display_value', None) or latest_call.stage.name

                last_conversation_outcome = (
                    latest_call.conversation_summary or 
                    call_tag or 
                    latest_call.connection_status or 
                    "-"
                )

            # Last Contacted Date & Time (Asia/Kolkata IST 12-hr)
            last_contacted_str = "-"
            if latest_call and latest_call.created_at:
                try:
                    last_contacted_str = timezone.localtime(latest_call.created_at).strftime("%d-%m-%Y %I:%M %p")
                except Exception:
                    last_contacted_str = latest_call.created_at.strftime("%d-%m-%Y %I:%M %p")

            # Inquiry Date & Time (Asia/Kolkata IST 12-hr)
            inquiry_dt = lead.enquiry_date or lead.created_at
            if inquiry_dt:
                try:
                    inquiry_date_str = timezone.localtime(inquiry_dt).strftime("%d-%m-%Y %I:%M %p")
                except Exception:
                    inquiry_date_str = inquiry_dt.strftime("%d-%m-%Y %I:%M %p")
            else:
                inquiry_date_str = "-"

            # Lead Age Calculation (e.g. "3 Months" / "15 Days")
            if inquiry_dt:
                delta_days = (timezone.now() - inquiry_dt).days
                if delta_days >= 30:
                    lead_age_str = f"{delta_days // 30} Months"
                else:
                    lead_age_str = f"{delta_days} Days"
            else:
                lead_age_str = "-"

            leads_list.append({
                "s_no": idx,
                "lead_id": lead.id,
                "name": lead.full_name or "",
                "contact": lead.mobile_no or "",
                "email": lead.email or "",
                "assigned_to": get_user_display_name(lead.assigned_to),
                "assigned_to_id": lead.assigned_to_id,
                "campaign_id": lead.campaign_id,
                "campaign_name": lead.campaign.name if lead.campaign else None,
                "effort_summary": effort_summary,
                "total_calls": total_calls,
                "loss_reason": loss_reason_str,
                "last_conversation_outcome": last_conversation_outcome,
                "last_contacted": last_contacted_str,
                "inquiry_date": inquiry_date_str,
                "lead_age": lead_age_str,
                "course": (
                    getattr(lead.course_name, 'course_name', None)
                    or getattr(getattr(lead.course, 'course_name', None), 'course_name', None)
                    or "-"
                ),
                "lead_source": lead.lead_source.name if lead.lead_source else "-",
                "approval_status": "pending_approval"
            })

        import math
        total_pages = math.ceil(total_count / page_size) if (page_size and page_size > 0 and total_count > 0) else 1

        return {
            "status": True,
            "message": "Loss lead approval requests fetched successfully!",
            "total_records": total_count,
            "total_count": total_count,
            "total_pages": total_pages,
            "page": page,
            "page_size": page_size,
            "showing_count": len(leads_list),
            "requests": leads_list,
            "leads": leads_list,
            "data": {
                "requests": leads_list,
                "leads": leads_list,
                "showing_count": len(leads_list),
                "total_records": total_count,
                "total_count": total_count,
                "total_pages": total_pages,
                "page": page,
                "page_size": page_size
            }
        }

    except Exception as e:
        raise APIException(str(e))


def export_loss_lead_approval_requests_admin(user=None, **data):
    try:
        data["page_size"] = "all"
        res = fetch_loss_lead_approval_requests_admin(user=user, **data)
        leads = res.get("requests") or res.get("leads") or res.get("data", {}).get("leads", [])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Loss_Lead_Approvals"

        from utils.excel_export import apply_excel_report_header, auto_fit_columns
        start_row = apply_excel_report_header(
            ws,
            report_name="Loss Lead Approvals Report",
            description="All in one loss lead approvals report",
            user=user
        )

        headers = [
            "S.No", "Name", "Contact", "Assigned to", "Effort Summary",
            "Loss Reason", "Last Conversation Outcome", "Last Contacted",
            "Inquiry Date", "Lead Age"
        ]

        header_fill = PatternFill(start_color="84C225", end_color="84C225", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        data_font = Font(name="Calibri", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        ws.row_dimensions[start_row].height = 26
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

        current_row = start_row + 1
        for idx, item in enumerate(leads, start=1):
            row_vals = [
                item.get("s_no", idx),
                item.get("name", "-"),
                item.get("contact", "-"),
                item.get("assigned_to", "-"),
                item.get("effort_summary", "-"),
                item.get("loss_reason", "-"),
                item.get("last_conversation_outcome", "-"),
                item.get("last_contacted", "-"),
                item.get("inquiry_date", "-"),
                item.get("lead_age", "-")
            ]
            ws.row_dimensions[current_row].height = 22
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                cell.alignment = center_align
            current_row += 1

        auto_fit_columns(ws, start_row=start_row, min_width=16, col_a_min=18, col_b_min=28)

        today_str = datetime.now().strftime("%Y%m%d")
        file_name = f"Loss_Lead_Approval_Requests_{today_str}.xlsx"
        export_dir = os.path.join(settings.MEDIA_ROOT, 'exports')
        os.makedirs(export_dir, exist_ok=True)
        file_path = os.path.join(export_dir, file_name)

        wb.save(file_path)

        return {
            "status": True,
            "message": f"Successfully exported {len(leads)} loss lead approval requests!",
            "total_exported": len(leads),
            "file_name": file_name,
            "download_url": f"/media/exports/{file_name}"
        }

    except Exception as e:
        raise APIException(str(e))


def action_loss_lead_approval_admin(user=None, **data):
    lead_id = data.get("lead_id") or data.get("id")
    if not lead_id:
        raise ValidationError("Lead ID is required.")

    lead = Lead.objects.filter(id=lead_id).first()
    if not lead:
        raise ValidationError("Lead not found.")

    raw_action = data.get("action") or data.get("action_type") or "approve"
    action_type = str(raw_action).lower().strip()

    admin_user = user if user and getattr(user, 'is_authenticated', False) else None
    if not admin_user:
        admin_user_id = data.get("user_id") or data.get("admin_user_id")
        if admin_user_id:
            admin_user = User.objects.filter(id=admin_user_id).first()
    
    if not admin_user:
        raise ValidationError("Authentication required. Please login to perform this action.")

    raw_remarks = (
        data.get("remarks") or 
        data.get("remark") or
        data.get("reason_for_rejection") or
        data.get("reasonForRejection") or
        data.get("rejection_reason") or
        data.get("rejectionReason") or
        data.get("rejection_remarks") or
        data.get("reject_reason") or
        data.get("reject_remarks") or
        data.get("final_remarks") or 
        data.get("reason") or 
        data.get("detailed_reason") or 
        data.get("notes") or 
        data.get("comments") or 
        data.get("description") or
        data.get("loss_remarks") or
        ""
    )
    remarks = str(raw_remarks).strip() if raw_remarks else ""
    if not remarks:
        known_keys = {'lead_id', 'id', 'action', 'action_type', 'assigned_to_id', 'telecaller_id', 'user_id', 'admin_user_id', 'can_retarget', 'page', 'page_size'}
        for k, v in data.items():
            if k not in known_keys and isinstance(v, str) and str(v).strip() != "":
                remarks = str(v).strip()
                break

    prev_assigned_to = lead.assigned_to

    with transaction.atomic():
        lead_pipe = (
            lead.pipeline or
            getattr(lead.pipeline_stage, 'pipeline_category', None) or
            (lead.campaign.pipeline_category if (lead.campaign and lead.campaign.pipeline_category) else None)
        )

        if action_type in ["approve", "accept", "confirm"]:
            loss_stage = None
            if lead_pipe:
                loss_stage = (
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, stage_type='closed_lost').first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, name__icontains="loss").first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, name__icontains="lost").first()
                )
            if not loss_stage and not lead_pipe:
                default_pipe = PipelineCategory.objects.first()
                if default_pipe:
                    loss_stage = PipelineStage.objects.filter(pipeline_category=default_pipe, stage_type='closed_lost').first()
                if not loss_stage:
                    loss_stage = PipelineStage.objects.filter(stage_type='closed_lost').first()

            if loss_stage:
                lead.pipeline_stage = loss_stage
                if not lead.pipeline and loss_stage.pipeline_category:
                    lead.pipeline = loss_stage.pipeline_category
            lead.current_status = "Loss"
            save_fields = ['pipeline_stage', 'current_status']
            if not lead.pipeline and lead.pipeline_id:
                save_fields.append('pipeline')
            lead.save(update_fields=save_fields)

            loss_detail = LossLeadDetail.objects.select_related("loss_reason").filter(lead=lead).first()
            main_reason_id_val = (
                loss_detail.loss_reason_id or loss_detail.main_reason_id
                if loss_detail else None
            )

            can_retarget_val = bool(data.get("can_retarget", True))
            created_by_user = get_user_display_name(admin_user) or "Admin"

            AdminApprovedLossLead.objects.update_or_create(
                lead=lead,
                defaults={
                    'approved_by': admin_user,
                    'main_reason_id': main_reason_id_val,
                    'final_remarks': remarks or (loss_detail.detailed_reason if loss_detail else "Loss Approved"),
                    'can_retarget': can_retarget_val,
                    'created_by': created_by_user,
                    'updated_by': created_by_user
                }
            )

            AdminLossActionLog.objects.create(
                lead=lead,
                admin_user=admin_user,
                action_type='approved',
                previous_assigned_to=prev_assigned_to,
                remarks=remarks or "Loss request approved",
                created_by=created_by_user,
                updated_by=created_by_user
            )

            result = {
                "message": f"Loss request for lead '{lead.full_name}' has been APPROVED successfully!",
                "lead_id": lead.id,
                "action": "approved"
            }

        elif action_type in ["reject", "deny", "cancel"]:
            followup_stage = None
            if lead_pipe:
                followup_stage = (
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, stage_type='follow_up').first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, name__icontains="follow").first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, stage_type="open", order_no__gt=1).first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, stage_type="open").first()
                )
            if not followup_stage and not lead_pipe:
                default_pipe = PipelineCategory.objects.first()
                if default_pipe:
                    followup_stage = PipelineStage.objects.filter(pipeline_category=default_pipe, stage_type='follow_up').first()
                if not followup_stage:
                    followup_stage = PipelineStage.objects.filter(stage_type='follow_up').first()

            if followup_stage:
                lead.pipeline_stage = followup_stage
                if not lead.pipeline and followup_stage.pipeline_category:
                    lead.pipeline = followup_stage.pipeline_category
            lead.current_status = "working"
            save_fields = ['pipeline_stage', 'current_status']
            if not lead.pipeline and lead.pipeline_id:
                save_fields.append('pipeline')
            lead.save(update_fields=save_fields)

            AdminApprovedLossLead.objects.filter(lead=lead).delete()

            created_by_user = get_user_display_name(admin_user) or "Admin"
            AdminLossActionLog.objects.create(
                lead=lead,
                admin_user=admin_user,
                action_type='rejected',
                previous_assigned_to=prev_assigned_to,
                remarks=remarks or "Loss request rejected",
                created_by=created_by_user,
                updated_by=created_by_user
            )

            result = {
                "message": f"Loss request for lead '{lead.full_name}' was REJECTED and restored to Follow Up!",
                "lead_id": lead.id,
                "action": "rejected"
            }

        elif action_type in ["reassign", "retry", "assign"]:
            raw_agent_val = (
                data.get("assigned_to_id")
                or data.get("new_telecaller_id")
                or data.get("telecaller_id")
                or data.get("user_id")
                or data.get("to_telecaller_id")
                or data.get("to_user_id")
                or data.get("new_agent_id")
                or data.get("telecaller")
            )
            if not raw_agent_val:
                raise ValidationError("Please select a telecaller to reassign this lead.")

            new_agent = None
            if str(raw_agent_val).isdigit():
                new_agent = User.objects.filter(id=int(raw_agent_val), is_active=True).first()
            if not new_agent:
                new_agent = User.objects.filter(
                    Q(username__iexact=str(raw_agent_val).strip()) |
                    Q(first_name__iexact=str(raw_agent_val).strip()),
                    is_active=True
                ).first()

            if not new_agent:
                raise ValidationError(f"Selected telecaller '{raw_agent_val}' not found or inactive.")

            if prev_assigned_to and prev_assigned_to.id == new_agent.id:
                raise ValidationError("Cannot reassign to the same telecaller. Please select a different telecaller.")

            lead.assigned_to = new_agent
            
            followup_stage = None
            if lead_pipe:
                followup_stage = (
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, stage_type='follow_up').first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, name__icontains="follow").first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, stage_type="open", order_no__gt=1).first() or
                    PipelineStage.objects.filter(pipeline_category=lead_pipe, stage_type="open").first()
                )
            if not followup_stage and not lead_pipe:
                default_pipe = PipelineCategory.objects.first()
                if default_pipe:
                    followup_stage = PipelineStage.objects.filter(pipeline_category=default_pipe, stage_type='follow_up').first()
                if not followup_stage:
                    followup_stage = PipelineStage.objects.filter(stage_type='follow_up').first()

            if followup_stage:
                lead.pipeline_stage = followup_stage
                if not lead.pipeline and followup_stage.pipeline_category:
                    lead.pipeline = followup_stage.pipeline_category
            lead.current_status = "working"
            save_fields = ['assigned_to', 'pipeline_stage', 'current_status']
            if not lead.pipeline and lead.pipeline_id:
                save_fields.append('pipeline')
            lead.save(update_fields=save_fields)

            AdminApprovedLossLead.objects.filter(lead=lead).delete()

            created_by_user = get_user_display_name(admin_user) or "Admin"
            AdminLossActionLog.objects.create(
                lead=lead,
                admin_user=admin_user,
                action_type='reassigned',
                previous_assigned_to=prev_assigned_to,
                new_assigned_to=new_agent,
                remarks=remarks or f"Reassigned to {get_user_display_name(new_agent)}",
                created_by=created_by_user,
                updated_by=created_by_user
            )
            # transaction.on_commit(lambda: send_lead_reassigned_notification(
            # lead.id, new_agent, prev_assigned_to.id if prev_assigned_to else None
            #     ))
            result = {
                "message": f"Lead '{lead.full_name}' has been REASSIGNED to {get_user_display_name(lead.assigned_to)}!",
                "lead_id": lead.id,
                "action": "reassigned",
                "assigned_to": get_user_display_name(lead.assigned_to)
            }
        else:
            raise ValidationError(f"Invalid action_type '{action_type}'. Valid actions: approve, reject, reassign.")

    # # ---------- Telecaller ku notification ----------
    # ---------- Notifications (transaction commit aana apram) ----------
    prev_id = prev_assigned_to.id if prev_assigned_to else None
    # new_agent = None
    if action_type in ["approve", "accept", "confirm"]:
        # Loss approve -> original telecaller ku
        if prev_id:
            send_loss_request_result_notification(
                lead.id, prev_id, True, admin_user.id, remarks or None
            )

    elif action_type in ["reject", "deny", "cancel"]:
        # Reject -> lead thirumba avanukke (assigned_to maaralai), avanukku result
        if prev_id:
            send_loss_request_result_notification(
                lead.id, prev_id, False, admin_user.id, remarks or None
            )

    elif action_type in ["reassign", "retry", "assign"] and new_agent:
        # Old telecaller ku: request reject + lead vera aalukku poiduchu
        if prev_id:
            send_loss_request_result_notification(
                lead.id, prev_id, False, admin_user.id,
                reassigned_to_id=new_agent.id,
            )
        # New telecaller ku: New Lead Assigned
        send_lead_assigned_notification(lead.id, new_agent.id, admin_user.id)




    Notification.objects.filter(
        notification_type='loss_request', lead_id=lead.id, is_read=False
    ).update(is_read=True)

    return result
