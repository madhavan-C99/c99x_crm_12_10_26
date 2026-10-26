import os
import csv
import math
import uuid
from datetime import datetime, date, timedelta
from django.conf import settings
from django.utils import timezone
from rest_framework.exceptions import APIException, ValidationError

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from telecalling.models import User
from telecalling.models.leads import CampaignName
from adm.models import Team
from adm.models.pipeline_category import PipelineCategory
from adm.models.export_log import ExportLog

from ..tasks.export_task import process_export_task, run_export_generation
from .lead_services import fetch_all_leads_admin
from .payment_services import fetch_all_pending_payments_admin
from .loss_lead_approval_services import fetch_loss_lead_approval_requests_admin
from .performance_services import fetch_performance_overview_admin
from .report_services import execute_report_service


FULL_ADMIN_ROLE_CODES = {"ADM", "MGR", "AMGR", "DEV", "SUPERADMIN", "ADMIN"}
TEAM_LEAD_ROLE_CODES = {"TEAM_LEAD", "TL"}

HEADER_DISPLAY_NAMES = {
    "s_no": "S.No",
    "full_name": "Name",
    "mobile_no": "Mobile No",
    "pipeline_name": "Pipeline",
    "pipeline_stage_name": "Stage",
    "assigned_to": "Assigned Telecaller",
    "priority_name": "Priority",
    "campaign_name": "Campaign",
    "source_name": "Source",
    "course_name": "Course Name",
    "course_plan_name": "Course Plan",
    "amount": "Amount",
    "pending_amount": "Pending Amount",
    "next_followup": "Next Followup",
    "last_contacted": "Last Contacted",
    "enquiry_date": "Enquiry Date",
    "user_name": "User Name",
    "reporting_manager": "Reporting Manager",
    "mobile_number": "Mobile Number",
    "date": "Date",
    "total_disposed_count": "Total Disposed Count",
    "disposed_yes_connected_count": "Disposed Yes Connected Count",
    "disposed_not_connected_count": "Disposed Not Connected Count",
    "total_followup_leads": "Total Followup Leads",
    "total_converted_leads": "Total Converted Leads",
    "total_lost_leads": "Total Lost Leads",
    "conversion_percentage": "Conversion %",
    "total_assigned_leads": "Total Assigned Leads",
    "followups_due_today": "Followups Due Today",
    "average_tat": "Average TAT",
    "campaign_name": "Campaign Name",
    "pipeline_name": "Pipeline Name",
    "total_lead_count": "Total Lead Count",
    "total_leads": "Total Leads",
    "total_unassigned_leads": "Total Unassigned Leads",
    "total_open_leads": "Total Open Leads",
    "total_in_progress_leads": "Total In-Progress Lead",
    "total_closed_leads": "Total Closed Leads",
    "total_disconnected_lost_leads": "Total Disconnected Lead (Lost)",
    "total_disconnected_pending_leads": "Total Disconnected (Pending)",
    "total_connected_converted_leads": "Total Connected Leads (Converted/Joined)",
    "total_connected_lost_leads": "Total Connected Leads (Lost/Closed)",
    "total_calls_connected": "Total Calls Connected",
    "leads_marked_as_converted": "Leads Marked As Converted",
    "leads_marked_as_lost": "Leads Marked As Lost",
    "date_range": "Date Range",
    "total_joined_converted": "Total Joined / Converted",
    "won_value": "Won Value",
    "total_lost_value": "Total Lost Value",
}


def get_user_role_code(user):
    if getattr(user, 'is_superuser', False):
        return "ADM"

    user_role_obj = user.user_roles.select_related('role').first() if hasattr(user, 'user_roles') else None
    if user_role_obj and user_role_obj.role and user_role_obj.role.code:
        return user_role_obj.role.code

    return ""


def is_admin_or_manager_user(user):
    role_code = get_user_role_code(user)
    return role_code in FULL_ADMIN_ROLE_CODES or role_code in TEAM_LEAD_ROLE_CODES


def get_team_user_ids(user):
    user_ids = {user.id}
    if hasattr(user, 'team') and user.team:
        member_ids = set(User.objects.filter(team=user.team).values_list('id', flat=True))
        user_ids.update(member_ids)
    
    led_team_ids = set(Team.objects.filter(leader=user).values_list('id', flat=True))
    if led_team_ids:
        led_member_ids = set(User.objects.filter(team_id__in=led_team_ids).values_list('id', flat=True))
        user_ids.update(led_member_ids)

    return list(user_ids)


def apply_role_data_scoping(user, entity, filters):
    role_code = get_user_role_code(user)

    if role_code in TEAM_LEAD_ROLE_CODES:
        key = "assigned_to_ids" if entity in ["leads", "pending_payments"] else "telecaller_ids"
        filters[key] = get_team_user_ids(user)
        if entity == "reports":
            filters["reporting_manager_id"] = user.id
    elif role_code not in FULL_ADMIN_ROLE_CODES:
        key = "assigned_to" if entity in ["leads", "pending_payments"] else "telecaller_id"
        filters[key] = user.id
        if entity == "reports":
            filters["user_id"] = user.id

    return filters


def fetch_export_records(entity, user, filters):
    entity = str(entity or "").strip().lower().replace("-", "_")
    if entity == "lead":
        entity = "leads"
    elif entity in ["pending_payment", "payments", "payment"]:
        entity = "pending_payments"
    elif entity in ["loss_approval", "loss", "loss_lead_approval", "loss_lead_approvals", "loss_leads"]:
        entity = "loss_approvals"
    elif entity in ["telecaller_performance", "performance_overview"]:
        entity = "performance"
    elif entity == "dashboard":
        entity = "telecalling_dashboard"

    filters = apply_role_data_scoping(user, entity, filters)

    # 🚀 1. Telecalling Dashboard (Single Entity)
    if entity == "telecalling_dashboard":
        from telecalling.services.dashboard_services import get_dashboard_pdf_data
        res = get_dashboard_pdf_data(user=user, **filters)
        return res.get("data", {})

    # 🚀 2. Daily Report
    if entity == "daily_report":
        from telecalling.services.file_services import download_daily_reports
        try:
            res = download_daily_reports(user=user, **filters)
            return res.get("data", {})
        except Exception:
            return {"user": getattr(user, 'username', ''), "report_date": filters.get("report_date", "")}

    # 🚀 3. Admin Dashboard (Single Entity)
    if entity == "admin_dashboard":
        from adm.services.admin_dashboard_services import fetch_admin_dashboard_export_data
        return fetch_admin_dashboard_export_data(user=user, **filters)

    if entity == "leads":
        filters["page_size"] = "all"
        return fetch_all_leads_admin(user=user, **filters).get("leads", [])

    if entity == "pending_payments":
        filters.pop("all_rows", None)
        filters["page"] = 1
        return fetch_all_pending_payments_admin(user=user, all_rows=True, **filters).get("leads", [])

    if entity in ["loss_approvals", "loss_lead_approval", "loss_lead_approvals"]:
        filters["page_size"] = "all"
        res = fetch_loss_lead_approval_requests_admin(user=user, **filters)
        if isinstance(res, dict):
            return res.get("requests") or res.get("leads") or res.get("data", {}).get("leads", [])
        return []

    if entity in ["performance", "performance_overview", "telecaller_performance"]:
        filters["page"] = 1
        res = fetch_performance_overview_admin(filters, user=user, all_rows=True)
        return res.get("performance_list") or res.get("data", {}).get("performance_list", [])

    if entity == "reports":
        report_payload = dict(filters or {})
        report_payload["page"] = 1
        report_payload["limit"] = 1000000
        res = execute_report_service(report_payload)
        data_res = res.get("data", [])
        if isinstance(data_res, dict) and "activity_logs" in data_res:
            return data_res.get("activity_logs", [])
        return data_res if isinstance(data_res, list) else []

    raise APIException(f"Unsupported export entity: {entity}")


COLUMN_SYNONYMS = {
    ('name', 'full_name', 'contact_name', 'student_name'): ['full_name', 'name', 'telecaller_name'],
    ('mobile', 'mobile_no', 'contact', 'phone', 'number'): ['mobile_no', 'contact', 'mobile'],
    ('stage', 'pipeline_stage', 'stage_name', 'pipeline_stage_name', 'status'): ['pipeline_stage_name', 'stage', 'approval_status', 'status'],
    ('pipeline', 'pipeline_name'): ['pipeline_name', 'pipeline'],
    ('priority', 'priority_name', 'tag'): ['priority_name', 'tag'],
    ('campaign', 'campaign_name'): ['campaign_name', 'campaign'],
    ('source', 'source_name', 'lead_source'): ['source_name', 'lead_source', 'source'],
    ('course', 'course_name'): ['course_name', 'course'],
    ('course_plan', 'course_plan_name', 'plan'): ['course_plan_name', 'course_plan'],
    ('next_follow_up', 'next_followup'): ['next_followup'],
    ('created', 'enquiry', 'joining', 'inquiry', 'date', 'enquiry_date'): ['enquiry_date', 'joining_date', 'inquiry_date', 'created_at'],
}


DEFAULT_UI_EXPORT_COLUMNS = {
    "leads": [
        "full_name",
        "mobile_no",
        "pipeline_name",
        "pipeline_stage_name",
        "assigned_to",
        "priority_name",
        "campaign_name",
        "source_name",
        "course_name",
        "course_plan_name",
        "next_followup",
        "amount",
        "pending_amount",
        "last_contacted",
        "enquiry_date"
    ],
    "pending_payments": [
        "full_name",
        "mobile_no",
        "course_name",
        "course_plan_name",
        "amount",
        "amount_paid",
        "pending_amount",
        "next_followup",
        "due_date",
        "payment_status",
        "enquiry_date"
    ],
    "loss_approvals": [
        "full_name",
        "mobile_no",
        "assigned_to",
        "pipeline_name",
        "loss_reason",
        "reported_by",
        "approval_status",
        "created_at"
    ],
    "loss_lead_approval": [
        "full_name",
        "mobile_no",
        "assigned_to",
        "pipeline_name",
        "loss_reason",
        "reported_by",
        "approval_status",
        "created_at"
    ],
    "loss_lead_approvals": [
        "full_name",
        "mobile_no",
        "assigned_to",
        "pipeline_name",
        "loss_reason",
        "reported_by",
        "approval_status",
        "created_at"
    ],
    "performance": [
        "user_name",
        "reporting_manager",
        "mobile_number",
        "total_assigned_leads",
        "total_calls_connected",
        "total_converted_leads",
        "total_lost_leads",
        "conversion_percentage"
    ]
}

INTERNAL_ID_COLUMNS = {
    "id", "pipeline_id", "pipeline_stage_id", "assigned_to_id",
    "priority_id", "campaign_id", "source_id", "course_name_id",
    "course_plan_id", "course_id", "user_id", "telecaller_id",
    "organization_id", "lead_id"
}

def is_internal_id(col):
    if not col:
        return False
    c = str(col).strip().lower()
    return c in INTERNAL_ID_COLUMNS or c.endswith("_id")


def resolve_export_columns(requested_cols, available_keys, entity=None):
    safe_available = [k for k in available_keys if not is_internal_id(k)]

    # 1. If no specific columns requested by UI, use clean default UI columns
    if not requested_cols:
        clean_entity = str(entity or "").strip().lower().replace("-", "_")
        if clean_entity in DEFAULT_UI_EXPORT_COLUMNS:
            default_cols = [c for c in DEFAULT_UI_EXPORT_COLUMNS[clean_entity] if c in available_keys]
            if default_cols:
                return default_cols
        return safe_available if safe_available else available_keys

    # 2. If frontend requested columns, clean them and filter out internal IDs
    clean_requested = [c for c in requested_cols if not is_internal_id(c)]
    if not clean_requested:
        clean_entity = str(entity or "").strip().lower().replace("-", "_")
        if clean_entity in DEFAULT_UI_EXPORT_COLUMNS:
            return [c for c in DEFAULT_UI_EXPORT_COLUMNS[clean_entity] if c in available_keys]
        return safe_available

    matched_cols = []
    for req in clean_requested:
        req_clean = str(req).strip().lower().replace(" ", "_")
        if is_internal_id(req_clean):
            continue
        if req_clean in available_keys:
            if req_clean not in matched_cols:
                matched_cols.append(req_clean)
            continue

        found = False
        for keys_tuple, target_list in COLUMN_SYNONYMS.items():
            if req_clean in keys_tuple:
                for target in target_list:
                    if target in available_keys and not is_internal_id(target) and target not in matched_cols:
                        matched_cols.append(target)
                        found = True
                        break
            if found:
                break

        if not found:
            for avail in safe_available:
                if req_clean in avail.lower() and avail not in matched_cols:
                    matched_cols.append(avail)
                    break

    final_cols = [c for c in matched_cols if not is_internal_id(c)]
    return final_cols if final_cols else safe_available


from utils.excel_export import apply_excel_report_header, auto_fit_columns, extract_row_value, format_generated_by


def generate_excel_file(file_path, export_cols, raw_rows, report_name="Report", description="All in one report", user=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export Data"

    # 1. Apply Top Metadata Header (Rows 1-4, spacers 5-6, returns row 7 for table headers)
    start_row = apply_excel_report_header(
        ws,
        report_name=report_name,
        description=description,
        user=user
    )

    headers = [HEADER_DISPLAY_NAMES.get(col, str(col).replace("_", " ").title()) for col in export_cols]

    header_fill = PatternFill(fill_type="solid", fgColor="84C225")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
    border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # 2. Style & Write Table Header Row (Row 7)
    ws.row_dimensions[start_row].height = 28
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    # 3. Write Data Rows (Row 8 onwards)
    current_row = start_row + 1
    for row in raw_rows:
        ws.row_dimensions[current_row].height = 24
        for col_idx, col in enumerate(export_cols, start=1):
            val = extract_row_value(row, col)
            if hasattr(val, "tzinfo") and getattr(val, "tzinfo", None) is not None:
                val = val.replace(tzinfo=None)
            elif isinstance(val, (datetime, date)):
                val = val.strftime("%d %b %Y, %I:%M %p") if isinstance(val, datetime) else val.strftime("%d %b %Y")
            elif val is None:
                val = ""
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = border
        current_row += 1

    # 4. Auto-Calculate Generous Column Widths
    auto_fit_columns(ws, start_row=start_row, min_width=16, col_a_min=18, col_b_min=28)

    ws.freeze_panes = f"A{start_row + 1}"
    if export_cols:
        max_col_letter = get_column_letter(len(export_cols))
        ws.auto_filter.ref = f"A{start_row}:{max_col_letter}{ws.max_row}"
    wb.save(file_path)


def export_data_service(user, entity, export_format, selected_columns, filters=None):
    try:
        filters = filters or {}

        raw_rows = fetch_export_records(entity, user, filters)
        print(f"Raw rows for export: {len(raw_rows) if isinstance(raw_rows, (list, dict)) else 0}")
        if not raw_rows:
            raw_rows = []

        if isinstance(raw_rows, list) and raw_rows:
            available_keys = list(raw_rows[0].keys())
        elif isinstance(raw_rows, dict):
            available_keys = list(raw_rows.keys())
        else:
            available_keys = []

        export_cols = resolve_export_columns(selected_columns, available_keys, entity=entity)
        if not export_cols and selected_columns:
            export_cols = [
                c.get("field") if isinstance(c, dict) else str(c).strip()
                for c in selected_columns
                if (isinstance(c, dict) and c.get("field")) or (isinstance(c, str) and str(c).strip())
            ]
        if not export_cols:
            export_cols = ["s_no"]

        media_export_dir = os.path.join(settings.MEDIA_ROOT, 'exports')
        os.makedirs(media_export_dir, exist_ok=True)
        report_key = str(filters.get("report_key") or "").strip()
        base_name = report_key if (entity == "reports" and report_key) else entity
        report_display = REPORT_DISPLAY_NAMES.get(report_key or entity, base_name.replace('_', ' ').title())
        desc = f"All in one {report_display.lower()}"
        ext = "csv" if export_format == "csv" else ("pdf" if export_format == "pdf" else "xlsx")
        today_str = datetime.now().strftime('%Y-%m-%d')
        display_prefix = report_display.replace(' ', '_')
        filename = f"{display_prefix}_{today_str}.{ext}"
        if os.path.exists(os.path.join(media_export_dir, filename)):
            filename = f"{display_prefix}_{today_str}_{uuid.uuid4().hex[:6]}.{ext}"
        file_path = os.path.join(media_export_dir, filename)

        if ext == "csv":
            generate_csv_file(
                file_path,
                export_cols,
                raw_rows,
                report_name=report_display,
                description=desc,
                user=user
            )
        elif ext == "pdf":
            from utils.pdf import SPECIALIZED_PDF_GENERATORS
            if entity in SPECIALIZED_PDF_GENERATORS:
                SPECIALIZED_PDF_GENERATORS[entity](file_path, raw_rows, user=user)
            elif report_key in SPECIALIZED_PDF_GENERATORS:
                SPECIALIZED_PDF_GENERATORS[report_key](file_path, raw_rows, user=user)
            else:
                generate_pdf_file(
                    file_path,
                    export_cols,
                    raw_rows,
                    report_title=report_display,
                    report_name=report_display,
                    description=desc,
                    user=user
                )
        else:
            if entity == "admin_dashboard":
                from adm.services.admin_dashboard_services import generate_admin_dashboard_excel
                generate_admin_dashboard_excel(file_path, raw_rows, user=user)
            else:
                generate_excel_file(
                    file_path,
                    export_cols,
                    raw_rows,
                    report_name=report_display,
                    description=desc,
                    user=user
                )

        total_count = len(raw_rows) if isinstance(raw_rows, list) else (1 if raw_rows else 0)
        custom_message = "Export generated successfully"

        return {
            "status": "success",
            "message": custom_message,
            "export_format": export_format,
            "total_exported": total_count,
            "file_name": filename,
            "download_url": f"{settings.MEDIA_URL}exports/{filename}"
        }

    except Exception as e:
        if isinstance(e, APIException):
            raise
        raise APIException(str(e)) from e


def generate_csv_file(file_path, export_cols, raw_rows, report_name="Report", description="All in one report", user=None):
    generated_by_str = format_generated_by(user)
    gen_date_str = timezone.localtime(timezone.now()).strftime("%a %b %d %Y")

    headers = [HEADER_DISPLAY_NAMES.get(col, str(col).replace("_", " ").title()) for col in export_cols]
    with open(file_path, mode='w', newline='', encoding='utf-8-sig') as f:
        writer = csv.writer(f)
        # 1. Top Metadata Header (Rows 1 to 4)
        writer.writerow(["Report Name", report_name])
        writer.writerow(["Description", description])
        writer.writerow(["Generated By", generated_by_str])
        writer.writerow(["Generated At", gen_date_str])
        writer.writerow([])
        writer.writerow([])

        # 2. Table Headers (Row 7)
        writer.writerow(headers)

        # 3. Data Rows with Smart Extractor (Row 8 onwards)
        for row in raw_rows:
            clean_row = []
            for col in export_cols:
                val = extract_row_value(row, col)
                if hasattr(val, "tzinfo") and getattr(val, "tzinfo", None) is not None:
                    val = val.replace(tzinfo=None)
                if isinstance(val, (datetime, date)):
                    val = val.strftime("%d %b %Y, %I:%M %p") if isinstance(val, datetime) else val.strftime("%d %b %Y")
                elif val is None:
                    val = ""
                clean_row.append(str(val))
            writer.writerow(clean_row)


def generate_pdf_file(file_path, export_cols, raw_rows, report_title="Export Report", report_name="Report", description="All in one report", user=None):
    doc = SimpleDocTemplate(
        file_path,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    meta_label_style = ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#000000')
    )
    meta_val_style = ParagraphStyle(
        'MetaVal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1F2937')
    )
    cell_style = ParagraphStyle(
        'CellStyle',
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1
    )
    header_style = ParagraphStyle(
        'HeaderStyle',
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1
    )

    generated_by_str = format_generated_by(user)
    gen_date_str = timezone.localtime(timezone.now()).strftime("%a %b %d %Y")
    display_report_name = report_name or report_title

    # 1. Top Metadata Header Table (Rows 1 to 4)
    meta_table_data = [
        [Paragraph("<b>Report Name</b>", meta_label_style), Paragraph(f": {display_report_name}", meta_val_style)],
        [Paragraph("<b>Description</b>", meta_label_style), Paragraph(f": {description}", meta_val_style)],
        [Paragraph("<b>Generated By</b>", meta_label_style), Paragraph(f": {generated_by_str}", meta_val_style)],
        [Paragraph("<b>Generated At</b>", meta_label_style), Paragraph(f": {gen_date_str}", meta_val_style)],
    ]
    meta_table = Table(meta_table_data, colWidths=[35 * mm, 150 * mm])
    meta_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
    ]))

    elements = [
        meta_table,
        Spacer(1, 6 * mm)
    ]

    # 2. Main Data Table with Smart Extractor
    headers = [Paragraph(HEADER_DISPLAY_NAMES.get(col, str(col).replace("_", " ").title()), header_style) for col in export_cols]
    table_data = [headers]

    for row in raw_rows:
        row_cells = []
        for col in export_cols:
            val = extract_row_value(row, col)
            if hasattr(val, "tzinfo") and getattr(val, "tzinfo", None) is not None:
                val = val.replace(tzinfo=None)
            if isinstance(val, (datetime, date)):
                val = val.strftime("%d %b %Y")
            elif val is None:
                val = ""
            row_cells.append(Paragraph(str(val), cell_style))
        table_data.append(row_cells)

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#84C225')),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t)
    doc.build(elements)


REPORT_DISPLAY_NAMES = {
    "leads": "Leads Report",
    "lead": "Leads Report",
    "pending_payments": "Pending Payments Report",
    "loss_approvals": "Loss Lead Approvals Report",
    "loss_lead_approval": "Loss Lead Approvals Report",
    "loss_lead_approvals": "Loss Lead Approvals Report",
    "performance": "Telecaller Performance Report",
    "admin_dashboard": "Admin Dashboard Report",
    "telecalling_dashboard": "Telecalling Dashboard Report",
    "daily_report": "Daily Telecaller Report",
    "user_report": "User Report",
    "lead_disposition_report": "Lead Disposition Report",
    "user_stage_report": "User Stage Report",
    "user_followup_report": "User Follow Up Report",
    "campaign_report": "Campaign Report",
    "campaign_lead_report": "Campaign Lead Report",
    "campaign_stage_report": "Campaign Stage Report",
    "user_deal_report": "User Deal Report",
    "login_report": "Login & Hourly Report",
    "hourly_login_report": "Login & Hourly Report"
}


def initiate_export_service(user, payload):
    try:
        report_key = str(payload.get("report_key") or payload.get("entity") or payload.get("export_type") or "").strip()
        if not report_key:
            raise ValidationError("report_key or entity is required")

        file_format = str(payload.get("file_format") or payload.get("format") or "xlsx").strip().lower()
        if file_format == "excel":
            file_format = "xlsx"
        if file_format not in ["xlsx", "csv", "pdf"]:
            file_format = "xlsx"

        raw_columns = payload.get("columns", [])
        columns = [
            c.get("field") if isinstance(c, dict) else str(c).strip()
            for c in raw_columns
            if (isinstance(c, dict) and c.get("field")) or (isinstance(c, str) and str(c).strip())
        ] if raw_columns else []

        filters = dict(payload.get("filters") or {})

        # Auto-pack any top-level filters if passed directly in payload
        for k, v in payload.items():
            if k not in ["report_key", "entity", "export_type", "report_type", "file_format", "format", "columns", "filters"] and v is not None:
                if k not in filters:
                    filters[k] = v

        report_type = REPORT_DISPLAY_NAMES.get(report_key, report_key.replace("_", " ").title() if report_key else "Reports")

        campaign_name = None
        pipeline_name = None
        if filters.get("campaign_id") and str(filters.get("campaign_id")).isdigit():
            c = CampaignName.objects.filter(id=int(filters["campaign_id"])).first()
            campaign_name = c.name if c else None

        if filters.get("pipeline_id") and str(filters.get("pipeline_id")).isdigit():
            p = PipelineCategory.objects.filter(id=int(filters["pipeline_id"])).first()
            pipeline_name = p.category_name if p else None

        # =========================================================================
        # 🚀 [ASYNC WORKER QUEUE - COMMENTED OUT AS PER SYNCHRONOUS DOWNLOAD REQUIREMENT]
        # Previous background worker queuing via Huey:
        # log = ExportLog.objects.create(
        #     user=user,
        #     report_key=report_key,
        #     report_type=report_type,
        #     campaign_name=campaign_name,
        #     pipeline_name=pipeline_name,
        #     file_format=file_format,
        #     status="pending"
        # )
        # # 🚀 Trigger Huey Background Task
        # process_export_task(log.id, user.id, report_key, file_format, columns, filters)
        # return {
        #     "status": "success",
        #     "message": "Export started",
        #     "data": {
        #         "export_id": log.id,
        #         "status": "pending"
        #     }
        # }
        # =========================================================================

        # 🚀 SYNCHRONOUS EXPORT GENERATION: Directly generate file (within 1-2s) & return URL immediately
        CORE_ENTITIES = [
            "leads",
            "pending_payments",
            "loss_approvals",
            "performance",
            "admin_dashboard",
            "telecalling_dashboard",
            "daily_report"
        ]

        clean_key = str(report_key or "").strip().lower().replace("-", "_")
        if clean_key == "lead":
            clean_key = "leads"
        elif clean_key in ["pending_payment", "payments", "payment"]:
            clean_key = "pending_payments"
        elif clean_key in ["loss_approval", "loss", "loss_lead_approval", "loss_lead_approvals", "loss_leads"]:
            clean_key = "loss_approvals"
        elif clean_key in ["telecaller_performance", "performance_overview"]:
            clean_key = "performance"
        elif clean_key == "dashboard":
            clean_key = "telecalling_dashboard"

        target_entity = clean_key if clean_key in CORE_ENTITIES else "reports"
        export_filters = dict(filters)
        export_filters["report_key"] = report_key

        export_res = export_data_service(
            user=user,
            entity=target_entity,
            export_format=file_format,
            selected_columns=columns,
            filters=export_filters
        )

        file_name = export_res.get("file_name") or ""
        download_url = export_res.get("download_url") or ""
        media_export_dir = os.path.join(settings.MEDIA_ROOT, 'exports')
        file_path = os.path.join(media_export_dir, file_name) if file_name else None

        # 🚀 Audit History in export_logs table with status "downloaded"
        log = ExportLog.objects.create(
            user=user,
            report_key=report_key,
            report_type=report_type,
            campaign_name=campaign_name,
            pipeline_name=pipeline_name,
            file_format=file_format,
            file_path=file_path,
            file_name=file_name,
            download_url=download_url,
            status="downloaded",
            expires_at=timezone.now() + timedelta(days=7)
        )

        return {
            "status": True,
            "message": "Export generated successfully",
            "file_name": file_name,
            "download_url": download_url,
            "data": {
                "export_id": log.id,
                "file_name": file_name,
                "download_url": download_url,
                "status": "downloaded"
            }
        }
    except ValidationError:
        raise
    except Exception as e:
        # 🚀 Audit failure in export_logs table
        try:
            ExportLog.objects.create(
                user=user,
                report_key=report_key,
                report_type=report_type,
                campaign_name=campaign_name,
                pipeline_name=pipeline_name,
                file_format=file_format,
                status="failed",
                error_message=str(e)
            )
        except Exception:
            pass
        raise APIException(str(e))


def fetch_export_logs_service(user, page=1, limit=25):
    qs = ExportLog.objects.all()
    if not is_admin_or_manager_user(user):
        qs = qs.filter(user=user)

    total_records = qs.count()
    try:
        limit_val = max(1, int(limit))
    except (ValueError, TypeError):
        limit_val = 25

    try:
        page_val = max(1, int(page))
    except (ValueError, TypeError):
        page_val = 1

    total_pages = math.ceil(total_records / limit_val) if total_records > 0 else 1
    start = (page_val - 1) * limit_val
    paged_logs = qs.select_related('user')[start:start + limit_val]

    logs_data = []
    for item in paged_logs:
        full_name = f"{item.user.first_name or ''} {item.user.last_name or ''}".strip() or item.user.username
        local_created = timezone.localtime(item.created_at) if item.created_at else None
        created_at_str = local_created.strftime('%Y-%m-%dT%H:%M:%S') if local_created else None
        logs_data.append({
            "id": item.id,
            "created_at": created_at_str,
            "requested_user_name": full_name,
            "report_type": item.report_type,
            "campaign_name": item.campaign_name,
            "pipeline_name": item.pipeline_name,
            "file_format": item.file_format,
            "file_name": item.file_name,
            "download_url": item.download_url,
            "status": item.status,
            "error_message": item.error_message
        })

    return {
        "data": logs_data,
        "total_records": total_records,
        "total_pages": total_pages,
        "current_page": page_val,
        "limit": limit_val
    }


def download_export_file_service(user, export_id):
    log = ExportLog.objects.filter(id=export_id).first()
    if not log:
        return {"_status": 404, "status": "error", "message": "Export ID not found"}

    if not is_admin_or_manager_user(user) and log.user_id != user.id:
        return {"_status": 403, "status": "error", "message": "Permission denied"}

    # 🚀 Self-Healing Fallback: If still pending or processing, generate on-demand so user is never stuck!
    if log.status in ['pending', 'processing']:
        if log.file_path and os.path.exists(log.file_path):
            log.status = "downloaded"
            log.save(update_fields=['status'])
        else:
            try:
                fallback_filters = {}
                if log.pipeline_name:
                    fallback_filters["pipeline_name"] = log.pipeline_name
                if log.campaign_name:
                    fallback_filters["campaign_name"] = log.campaign_name

                run_export_generation(
                    export_id=log.id,
                    user_id=log.user_id,
                    report_key=log.report_key,
                    file_format=log.file_format or "xlsx",
                    columns=[],
                    filters=fallback_filters
                )
                log.refresh_from_db()
            except Exception as e:
                log.status = "failed"
                log.error_message = str(e)
                log.save(update_fields=['status', 'error_message'])

    if log.status in ['pending', 'processing']:
        return {"_status": 409, "status": "processing", "message": "Export is still being generated. Please wait."}

    if log.status == 'failed':
        return {"_status": 400, "status": "failed", "message": f"Export failed: {log.error_message}"}

    if log.expires_at and timezone.now() > log.expires_at:
        return {"_status": 410, "status": "error", "message": "Export file has expired"}

    # Mark status as downloaded if re-downloaded
    if log.status != 'downloaded':
        log.status = 'downloaded'
        log.save(update_fields=['status'])

    return {
        "_status": 200,
        "status": "success",
        "data": {
            "export_id": log.id,
            "file_name": log.file_name,
            "download_url": log.download_url
        }
    }
